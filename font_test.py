from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

fonts = ['NSimSun', 'Songti SC', 'STSong', 'Hiragino Sans GB', 'Hiragino Sans GB W3', 'Hiragino Sans GB W6', 'PingFang SC', 'FangSong', 'KaiTi']
doc=Document()
for name in fonts:
    p=doc.add_paragraph(); r=p.add_run(name+'：中文测试 DevPilot 智能研发系统')
    r.font.name=name; r._element.rPr.rFonts.set(qn('w:ascii'),name); r._element.rPr.rFonts.set(qn('w:hAnsi'),name); r._element.rPr.rFonts.set(qn('w:eastAsia'),name); r.font.size=Pt(16)
doc.save('tmp/font_test.docx')
