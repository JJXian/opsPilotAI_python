import ast
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).parent
OUT = ROOT / 'output' / 'docx'
OUT.mkdir(parents=True, exist_ok=True)
TARGET = OUT / 'DevPilot智能研发系统_项目面试准备手册.docx'

BLUE = '12355B'; TEAL = '0B6E8A'; INK = '202A35'; MUTED = '657786'; GOLD = '7C3E00'

CN_BODY = 'PingFang SC'
CN_HEAD = 'PingFang SC'

def font(run, name=CN_BODY, size=10.5, color=INK, bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:ascii'), name)
    run._element.rPr.rFonts.set(qn('w:hAnsi'), name)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold

def set_para(p, before=0, after=6, line=1.25, keep=False):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    p.paragraph_format.keep_with_next = keep

def add_text(doc, text, kind='body'):
    p = doc.add_paragraph()
    if kind == 'h1':
        set_para(p, 18, 10, 1.2, True); font(p.add_run(text), CN_HEAD, 16, BLUE, True)
    elif kind == 'h2':
        set_para(p, 12, 7, 1.2, True); font(p.add_run(text), CN_HEAD, 13, TEAL, True)
    elif kind == 'q':
        set_para(p, 10, 4, 1.25, True); font(p.add_run(text), CN_HEAD, 12, '152D4F', True)
    elif kind == 'note':
        set_para(p, 2, 4, 1.2); font(p.add_run(text), CN_BODY, 9, MUTED)
    elif kind == 'label':
        set_para(p, 3, 2, 1.25); font(p.add_run(text), CN_HEAD, 9.5, GOLD, True)
    elif kind == 'bullet':
        set_para(p, 0, 4, 1.25); p.paragraph_format.left_indent=Inches(.38); p.paragraph_format.first_line_indent=Inches(-.19); font(p.add_run(text), CN_BODY, 10.5)
    else:
        set_para(p, 0, 6, 1.25); font(p.add_run(text), CN_BODY, 10.5)
    return p

def add_labeled(doc, label, text):
    p=doc.add_paragraph(); set_para(p, 1, 5, 1.25)
    font(p.add_run(label), CN_HEAD, 10.5, '7C3E00', True)
    font(p.add_run(text), CN_BODY, 10.5, INK)
    return p

def page_number(paragraph):
    paragraph.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=paragraph.add_run('第 '); font(r,CN_BODY,8.5,MUTED)
    fldChar1 = OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText'); instrText.set(qn('xml:space'), 'preserve'); instrText.text='PAGE'
    fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'), 'end')
    r._r.append(fldChar1); r._r.append(instrText); r._r.append(fldChar2)
    font(paragraph.add_run(' 页'),CN_BODY,8.5,MUTED)

def setup(doc):
    sec=doc.sections[0]
    sec.top_margin=Inches(.78); sec.bottom_margin=Inches(.72); sec.left_margin=Inches(.82); sec.right_margin=Inches(.82)
    sec.header_distance=Inches(.32); sec.footer_distance=Inches(.34)
    head=sec.header.paragraphs[0]; head.alignment=WD_ALIGN_PARAGRAPH.LEFT
    font(head.add_run('DevPilot 智能研发系统｜项目面试准备手册'),CN_BODY,8.5,MUTED)
    foot=sec.footer.paragraphs[0]; page_number(foot)
    st=doc.styles['Normal']; st.font.name=CN_BODY; st._element.rPr.rFonts.set(qn('w:ascii'),CN_BODY); st._element.rPr.rFonts.set(qn('w:hAnsi'),CN_BODY); st._element.rPr.rFonts.set(qn('w:eastAsia'),CN_BODY); st.font.size=Pt(10.5)

def replacement(s):
    return (s.replace('Context Recall、Hit Rate 都是 0.988，MRR 是 0.935', '在 Top-3 检索口径下，Context Recall 为 0.918、Hit Rate 为 0.885、MRR 为 0.823')
             .replace('Context Recall 0.988、Hit Rate 0.988、MRR 0.935', 'Context Recall 0.918、Hit Rate 0.885、MRR 0.823')
             .replace('0.988 和 0.935', '0.918、0.885 和 0.823'))

def get_calls():
    tree=ast.parse((ROOT/'build_interview_pdf.py').read_text())
    result=[]
    for node in tree.body:
        if not isinstance(node, ast.AugAssign): continue
        val=node.value
        calls=[]
        if isinstance(val, ast.Call): calls=[val]
        elif isinstance(val, (ast.List,ast.Tuple)):
            calls=[x for x in val.elts if isinstance(x,ast.Call)]
        for c in calls:
            if not isinstance(c.func, ast.Name): continue
            name=c.func.id
            if name not in {'section','question','P'}: continue
            args=[]
            try: args=[ast.literal_eval(x) for x in c.args]
            except Exception: continue
            result.append((name,args))
    return result

def main():
    doc=Document(); setup(doc)
    # Cover: editorial_cover pattern adapted for a compact reference guide.
    for _ in range(5): doc.add_paragraph()
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_para(p,0,16,1.1)
    font(p.add_run('DevPilot 智能研发系统'),CN_HEAD,25,BLUE,True)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_para(p,0,14,1.2)
    font(p.add_run('项目面试准备手册｜Java 后端 / 全栈 / AI 应用 / Agent 开发'),CN_BODY,12,'4B6178')
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_para(p,0,22,1.2)
    font(p.add_run('候选人背景：硕士 · 约 1 年工作经验 · 目标：杭州中大厂'),CN_BODY,11,'4B6178')
    add_text(doc,'使用方式：先背“90 秒项目介绍”，再按岗位重点复习题库。所有带“建议补充”的地方，请用你的真实实现替换；不要把未知数据当作既成事实。','note')
    doc.add_page_break()
    add_text(doc,'目录','h1')
    for s in ['1. 项目定位与 90 秒开场','2. 项目全链路拆解','3. 核心题库：Agent 与 RAG','4. 核心题库：检索、评测与数据','5. 核心题库：工程、架构与可靠性','6. 核心题库：Bugfix Agent','7. 数据库、高并发与 Java 迁移','8. 弱点、质疑点与复习清单']:
        add_text(doc,s,'bullet')
    doc.add_page_break()
    for name,args in get_calls():
        if name=='section':
            add_text(doc,replacement(args[0]),'h1')
            if len(args)>1 and args[1]: add_text(doc,replacement(args[1]),'body')
        elif name=='question':
            n,title,answer=args[:3]
            add_text(doc, f'{n}. {replacement(title)}','q')
            add_labeled(doc,'【参考回答】',replacement(answer))
            if len(args)>3 and args[3]: add_labeled(doc,'【建议补充】',replacement(args[3]))
            if len(args)>4 and args[4]: add_labeled(doc,'【通用原理】',replacement(args[4]))
        elif name=='P':
            text=args[0]
            if text.startswith('<b>'):
                # Keep lead paragraphs legible; strip the very small HTML subset used by the PDF builder.
                import re
                text=re.sub(r'</?b>','',text).replace('<br/>','')
            if text.startswith('• '): add_text(doc,replacement(text),'bullet')
            elif text not in ('目录',): add_text(doc,replacement(text),'body')
    doc.save(TARGET)
    print(TARGET)

if __name__=='__main__': main()
