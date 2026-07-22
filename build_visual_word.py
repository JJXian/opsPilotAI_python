from pathlib import Path
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

root=Path(__file__).parent
out=root/'output'/'docx'
out.mkdir(parents=True, exist_ok=True)
target=out/'DevPilot智能研发系统_项目面试准备手册_版式保留版.docx'
pages=sorted((root/'tmp'/'word_pages').glob('page-*.png'))
doc=Document()
sec=doc.sections[0]
sec.top_margin=Inches(.15); sec.bottom_margin=Inches(.15); sec.left_margin=Inches(.15); sec.right_margin=Inches(.15)
for i, image in enumerate(pages):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=0; p.paragraph_format.space_after=0
    p.add_run().add_picture(str(image), width=Inches(7.30))
    if i != len(pages)-1: doc.add_page_break()
doc.save(target)
print(target)
