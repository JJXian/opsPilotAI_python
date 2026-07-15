from docx import Document as DocxDocument

from app.services.document_splitter_service import document_splitter_service


def test_split_docx_preserves_headings_paragraphs_and_tables(tmp_path):
    file_path = tmp_path / "runbook.docx"
    document = DocxDocument()
    document.add_heading("服务故障处理手册", level=1)
    document.add_paragraph("遇到连接超时时，先检查应用日志和数据库连接池。")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "告警名称"
    table.rows[0].cells[1].text = "处理动作"
    table.rows[1].cells[0].text = "数据库连接超时"
    table.rows[1].cells[1].text = "检查连接池"
    document.save(file_path)

    documents = document_splitter_service.split_docx(str(file_path))

    content = "\n".join(item.page_content for item in documents)
    assert "服务故障处理手册" in content
    assert "应用日志" in content
    assert "告警名称：数据库连接超时" in content
    assert "处理动作：检查连接池" in content
    assert all(item.metadata["_extension"] == ".docx" for item in documents)
    assert all(item.metadata["_file_name"] == "runbook.docx" for item in documents)
